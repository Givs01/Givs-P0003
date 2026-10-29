"""
Contents (Index) shared by the Excel output and the PDF report.

Columns: Sr. No. | City Name | Page No. | Remarks
Page numbers refer to the PDF report: the contents page(s) come first,
then one A4 page per city.
"""
import math

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HEADERS = ["Sr. No.", "City Name", "Page No.", "Remarks"]
ZERO_REMARK = "All values are 0"

# ---------------------------------------------------------
# PDF LAYOUT (points, A4 portrait) - also decides page numbers
# ---------------------------------------------------------
PAGE_MARGIN_X = 40
PAGE_TOP = 50
TITLE_H = 44
HEADER_H = 24
ROW_H = 20
PAGE_BOTTOM = 60
A4_H = 841.89

ROWS_PER_PAGE = int((A4_H - PAGE_TOP - TITLE_H - HEADER_H - PAGE_BOTTOM) // ROW_H)


def contents_page_count(n_cities):
    return max(1, math.ceil(n_cities / ROWS_PER_PAGE))


def city_page_numbers(n_cities):
    first = contents_page_count(n_cities) + 1
    return [first + i for i in range(n_cities)]


# ---------------------------------------------------------
# EXCEL INDEX SHEET
# ---------------------------------------------------------
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=10)
LINK_FONT = Font(name="Arial", size=10, color="0563C1", underline="single")
STRIPE_FILL = PatternFill("solid", fgColor="F2F5F9")
ZERO_FILL = PatternFill("solid", fgColor="FFC7CE")
ZERO_FONT = Font(name="Arial", size=10, bold=True, color="9C0006")
LINE = Side(style="thin", color="BFC7D1")
BOX = Border(left=LINE, right=LINE, top=LINE, bottom=LINE)


def write_index(ws, entries):
    """
    ws: the (empty) Index worksheet.
    entries: list of (sheet_name, all_zero) in merge order.
    """
    pages = city_page_numbers(len(entries))

    for col, (title, width) in enumerate(zip(HEADERS, [9, 45, 11, 24]), start=1):
        cell = ws.cell(1, col, title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = BOX
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[cell.column_letter].width = width
    ws.row_dimensions[1].height = 22

    for i, ((name, all_zero), page) in enumerate(zip(entries, pages)):
        r = i + 2
        values = [i + 1, name, page, ZERO_REMARK if all_zero else None]

        for col, value in enumerate(values, start=1):
            cell = ws.cell(r, col, value)
            cell.font = BODY_FONT
            cell.border = BOX
            cell.alignment = Alignment(
                horizontal="left" if col in (2, 4) else "center",
                vertical="center",
            )
            if all_zero:
                cell.fill = ZERO_FILL
            elif i % 2:
                cell.fill = STRIPE_FILL

        link = ws.cell(r, 2)
        link.hyperlink = f"#'{name}'!A1"
        link.font = LINK_FONT
        if all_zero:
            ws.cell(r, 4).font = ZERO_FONT
        ws.row_dimensions[r].height = 18

    ws.freeze_panes = "A2"
    ws.print_title_rows = "1:1"

    # A4 PAGE SETTINGS
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True

    ws.page_margins.left = 0.4
    ws.page_margins.right = 0.4
    ws.page_margins.top = 0.75
    ws.page_margins.bottom = 0.75


def read_index(ws):
    """(city name, remark) rows back from a written Index sheet."""
    rows = []
    for r in range(2, ws.max_row + 1):
        name = ws.cell(r, 2).value
        if name:
            rows.append((str(name), ws.cell(r, 4).value or ""))
    return rows

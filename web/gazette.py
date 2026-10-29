"""
Browser version of main.py (runs inside Pyodide).

Reads .xls files directly with xlrd (no Excel / win32com needed),
copies values + formatting into one openpyxl workbook and applies
the same post-processing as the desktop script.
"""
from io import BytesIO

import xlrd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from contents import city_page_numbers, write_index
from pdf_report import build_pdf

# =========================================================
# CONFIG (same as main.py)
# =========================================================
DELETE_START_ROW = 40
DELETE_ROW_COUNT = 5
DEFAULT_EXCEL_WIDTH = 17.6

OUTPUT_NAME = "SLB Gazette file_merged.xlsx"
PDF_NAME = "SLB Gazette file_merged.pdf"

# xlrd code -> openpyxl name
BORDER_STYLES = {
    1: "thin", 2: "medium", 3: "dashed", 4: "dotted", 5: "thick",
    6: "double", 7: "hair", 8: "mediumDashed", 9: "dashDot",
    10: "mediumDashDot", 11: "dashDotDot", 12: "mediumDashDotDot",
    13: "slantDashDot",
}
H_ALIGN = {
    1: "left", 2: "center", 3: "right", 4: "fill", 5: "justify",
    6: "centerContinuous", 7: "distributed",
}
V_ALIGN = {0: "top", 1: "center", 2: "bottom", 3: "justify", 4: "distributed"}


# =========================================================
# XLS STYLE CONVERSION
# =========================================================
def _rgb(book, colour_index):
    rgb = book.colour_map.get(colour_index)
    if not rgb:
        return None
    return "FF%02X%02X%02X" % rgb


def _side(book, line_style, colour_index):
    style = BORDER_STYLES.get(line_style)
    if not style:
        return Side()
    return Side(style=style, color=_rgb(book, colour_index))


def _convert_xf(book, xf_index):
    xf = book.xf_list[xf_index]

    f = book.font_list[xf.font_index]
    font = Font(
        name=f.name,
        size=f.height / 20,
        # many non-Excel writers set only the weight (700), not the bold flag
        bold=bool(f.bold) or f.weight >= 600,
        italic=bool(f.italic),
        underline="single" if f.underline_type else None,
        strike=bool(f.struck_out),
        color=_rgb(book, f.colour_index),
    )

    bg = xf.background
    if bg.fill_pattern == 1:
        colour = _rgb(book, bg.pattern_colour_index)
        fill = PatternFill("solid", fgColor=colour) if colour else PatternFill()
    else:
        fill = PatternFill()

    b = xf.border
    border = Border(
        left=_side(book, b.left_line_style, b.left_colour_index),
        right=_side(book, b.right_line_style, b.right_colour_index),
        top=_side(book, b.top_line_style, b.top_colour_index),
        bottom=_side(book, b.bottom_line_style, b.bottom_colour_index),
    )

    a = xf.alignment
    rotation = a.rotation if 0 <= a.rotation <= 180 else 0
    alignment = Alignment(
        horizontal=H_ALIGN.get(a.hor_align),
        vertical=V_ALIGN.get(a.vert_align),
        wrap_text=bool(a.text_wrapped),
        text_rotation=rotation,
        indent=a.indent_level,
    )

    fmt = book.format_map.get(xf.format_key)
    number_format = fmt.format_str if fmt else "General"

    return font, fill, border, alignment, number_format


def _cell_value(cell):
    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        return None
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    if cell.ctype == xlrd.XL_CELL_ERROR:
        return xlrd.error_text_from_code.get(cell.value)
    if cell.ctype == xlrd.XL_CELL_NUMBER and float(cell.value).is_integer():
        return int(cell.value)
    if isinstance(cell.value, str):
        # \r\n would turn into a blank line once saved as .xlsx
        return cell.value.replace("\r\n", "\n").replace("\r", "\n")
    return cell.value


def _copy_sheet(book, source_ws, target_ws, style_cache):

    # -----------------------------------------------------
    # COPY CELL VALUES + STYLES
    # -----------------------------------------------------
    for r in range(source_ws.nrows):
        for c in range(source_ws.ncols):

            cell = source_ws.cell(r, c)
            value = _cell_value(cell)
            xf_index = cell.xf_index

            if value is None and xf_index is None:
                continue

            new_cell = target_ws.cell(row=r + 1, column=c + 1)
            new_cell.value = value

            if xf_index is not None:
                if xf_index not in style_cache:
                    style_cache[xf_index] = _convert_xf(book, xf_index)

                font, fill, border, alignment, number_format = style_cache[xf_index]
                new_cell.font = font
                new_cell.fill = fill
                new_cell.border = border
                new_cell.alignment = alignment
                new_cell.number_format = number_format

    # -----------------------------------------------------
    # COPY COLUMN WIDTHS
    # -----------------------------------------------------
    for col_idx in range(source_ws.ncols):

        col_letter = get_column_letter(col_idx + 1)

        info = source_ws.colinfo_map.get(col_idx)

        if info and info.width:
            target_ws.column_dimensions[col_letter].width = info.width / 256
        else:
            target_ws.column_dimensions[col_letter].width = DEFAULT_EXCEL_WIDTH

    # -----------------------------------------------------
    # FORCE D & E SAME AS C
    # -----------------------------------------------------
    c_width = target_ws.column_dimensions["C"].width or DEFAULT_EXCEL_WIDTH

    target_ws.column_dimensions["C"].width = c_width
    target_ws.column_dimensions["D"].width = c_width
    target_ws.column_dimensions["E"].width = c_width

    # -----------------------------------------------------
    # COPY ROW HEIGHTS
    # -----------------------------------------------------
    # Only rows with a custom height. Automatic rows are left for Excel to
    # size, exactly like the Excel-converted file in main.py.
    target_ws.sheet_format.defaultRowHeight = source_ws.default_row_height / 20

    for r, info in source_ws.rowinfo_map.items():
        if info.hidden:
            target_ws.row_dimensions[r + 1].hidden = True
        if info.height_mismatch:
            target_ws.row_dimensions[r + 1].height = info.height / 20

    # -----------------------------------------------------
    # COPY MERGED CELLS
    # -----------------------------------------------------
    for rlo, rhi, clo, chi in source_ws.merged_cells:
        target_ws.merge_cells(
            start_row=rlo + 1, start_column=clo + 1,
            end_row=rhi, end_column=chi,
        )


def _to_bytes(data):
    # JS Uint8Array arrives as a memoryview or a JsProxy
    if hasattr(data, "to_bytes"):
        return data.to_bytes()
    return bytes(data)


# =========================================================
# ALL-ZERO CHECK
# =========================================================
def is_all_zero(ws):
    """True when the sheet has numbers and every one of them is 0."""
    found_number = False

    for row in ws.iter_rows(values_only=True):
        for value in row:
            if isinstance(value, bool) or not isinstance(value, (int, float)):
                continue
            if value != 0:
                return False
            found_number = True

    return found_number


# =========================================================
# MAIN ENTRY (called from app.js)
# =========================================================
def build(files, font_dir="fonts"):
    """
    files: list of (file_name, bytes) in the required merge order.
    Returns (xlsx_bytes, pdf_bytes, report, log); report is a list of dicts per sheet.
    """
    log = []

    merged_wb = Workbook()
    merged_wb.remove(merged_wb.active)

    # -----------------------------------------------------
    # CREATE INDEX SHEET
    # -----------------------------------------------------
    # Filled in after all sheets are added (see write_index below)
    index_ws = merged_wb.create_sheet(title="Index", index=0)

    entries = []
    report = []

    # -----------------------------------------------------
    # PROCESS EACH FILE
    # -----------------------------------------------------
    for file_name, data in files:

        log.append(f"Processing: {file_name}")

        try:
            book = xlrd.open_workbook(
                file_contents=_to_bytes(data), formatting_info=True
            )
        except Exception as e:
            raise ValueError(f"Could not read '{file_name}': {e}") from e

        style_cache = {}

        for source_ws in book.sheets():

            sheet_name = source_ws.name

            # UNIQUE SHEET NAME
            new_sheet_name = sheet_name
            counter = 1

            while new_sheet_name in merged_wb.sheetnames:
                new_sheet_name = f"{sheet_name}_{counter}"
                counter += 1

            target_ws = merged_wb.create_sheet(title=new_sheet_name)

            _copy_sheet(book, source_ws, target_ws, style_cache)

            # DELETE ROWS 40-44
            delete_end_row = DELETE_START_ROW + DELETE_ROW_COUNT - 1

            for merged in list(target_ws.merged_cells.ranges):
                if (
                    merged.min_row <= delete_end_row
                    and merged.max_row >= DELETE_START_ROW
                ):
                    target_ws.unmerge_cells(str(merged))

            target_ws.delete_rows(DELETE_START_ROW, DELETE_ROW_COUNT)

            for r in range(DELETE_START_ROW, DELETE_START_ROW + DELETE_ROW_COUNT):
                if r in target_ws.row_dimensions:
                    del target_ws.row_dimensions[r]

            # INDEX ENTRY
            all_zero = is_all_zero(target_ws)
            entries.append((new_sheet_name, all_zero))

            report.append({
                "sr": len(entries),
                "sheet": new_sheet_name,
                "file": file_name,
                "all_zero": all_zero,
            })

            log.append(f"  Added Sheet: {new_sheet_name}"
                       + ("  [ALL ZERO]" if all_zero else ""))

    # INDEX: Sr. No. | City Name | Page No. | Remarks
    write_index(index_ws, entries)

    pages = city_page_numbers(len(entries))
    for row, page in zip(report, pages):
        row["page"] = page

    # -----------------------------------------------------
    # SAVE + REOPEN FOR FINAL PROCESSING (mirrors main.py)
    # -----------------------------------------------------
    buffer = BytesIO()
    merged_wb.save(buffer)
    buffer.seek(0)

    wb = load_workbook(buffer)

    for ws in wb.worksheets:

        if ws.title.strip().lower() == "index":
            continue

        # FIND LAST NON-EMPTY ROW
        last_row = ws.max_row

        while last_row > 1:
            row_has_data = any(
                ws.cell(last_row, col).value is not None
                for col in range(1, ws.max_column + 1)
            )
            if row_has_data:
                break
            last_row -= 1

        # MERGE A:E IN LAST ROW
        ws.merge_cells(f"A{last_row}:E{last_row}")

        ws[f"A{last_row}"].alignment = Alignment(
            horizontal="left",
            vertical="center"
        )

        # DELETE EXTRA ROWS AFTER LAST ROW
        rows_after = ws.max_row - last_row

        if rows_after > 0:
            ws.delete_rows(last_row + 1, rows_after)

        # PAGE SETUP
        ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
        ws.page_setup.paperSize = ws.PAPERSIZE_A4

        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True

        ws.page_margins.left = 0.25
        ws.page_margins.right = 0.25
        ws.page_margins.top = 0.4
        ws.page_margins.bottom = 0.4
        ws.page_margins.header = 0.2
        ws.page_margins.footer = 0.2

        ws.print_options.horizontalCentered = True
        ws.print_options.verticalCentered = False

        ws.print_title_rows = "1:1"

        last_col_letter = get_column_letter(ws.max_column)
        ws.print_area = f"A1:{last_col_letter}{last_row}"

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True

    out = BytesIO()
    wb.save(out)

    # PDF REPORT: Index first, then one A4 page per city sheet
    pdf_bytes = build_pdf(wb, font_dir)
    log.append(f"PDF: {len(wb.worksheets) - 1} city page(s) + Index")

    return out.getvalue(), pdf_bytes, report, log

"""
PDF report for the merged Gazette workbook (runs inside Pyodide).

- Index sheet first (continues over pages if long, header repeated)
- Every other sheet (one city) scaled to fit exactly ONE A4 portrait page
- Index entries are clickable links to the city page
"""
import datetime
import os
import re
import unicodedata
from io import BytesIO

import contents
from fontTools.ttLib import TTFont
from fpdf import FPDF
from openpyxl.styles.numbers import is_date_format
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel

A4_W, A4_H = 595.28, 841.89          # points
DEFAULT_COL_WIDTH = 8.43             # Excel characters
DEFAULT_ROW_HEIGHT = 15              # points
PAD = 2                              # cell padding (points, before scaling)

BORDER_WIDTHS = {
    "hair": 0.25, "thin": 0.5, "dotted": 0.5, "dashed": 0.5,
    "dashDot": 0.5, "dashDotDot": 0.5, "slantDashDot": 1.0,
    "medium": 1.0, "mediumDashed": 1.0, "mediumDashDot": 1.0,
    "mediumDashDotDot": 1.0, "double": 1.5, "thick": 1.5,
}
DASHES = {
    "dotted": (0.5, 1), "dashed": (3, 1.5), "mediumDashed": (3, 1.5),
    "dashDot": (3, 1.5), "dashDotDot": (3, 1.5),
    "mediumDashDot": (3, 1.5), "mediumDashDotDot": (3, 1.5),
}
THEME_COLOURS = {10: (5, 99, 193)}   # theme 10 = hyperlink blue

LINE_RATIO = 1.3                     # line height / font size (as Excel autofit)

# Bundled fonts (web/fonts). Arimo/Carlito/Tinos have exactly the same
# character widths as Arial/Calibri/Times New Roman, so text takes the
# same space as in Excel.
BUNDLED = {
    "Arimo": ("Arimo-Regular.ttf", "Arimo-Bold.ttf"),
    "Carlito": ("Carlito-Regular.ttf", "Carlito-Bold.ttf"),
    "Tinos": ("Tinos-Regular.ttf", "Tinos-Bold.ttf"),
    "DejaVu": ("DejaVuSans.ttf", "DejaVuSans-Bold.ttf"),
    "NotoGujarati": ("NotoSansGujarati-Regular.ttf", "NotoSansGujarati-Bold.ttf"),
    "NotoDevanagari": ("NotoSansDevanagari-Regular.ttf", "NotoSansDevanagari-Bold.ttf"),
}
METRIC_FONTS = {
    "arial": "Arimo", "helvetica": "Arimo",
    "calibri": "Carlito",
    "times new roman": "Tinos", "times": "Tinos",
}
DEFAULT_FAMILY = "Arimo"

# Tried in order for characters the cell's font does not have
FALLBACKS = ["DejaVu", "NotoGujarati", "NotoDevanagari"]
COMPLEX = {"NotoGujarati", "NotoDevanagari"}     # need shaping


def _system_font_files(name):
    """Windows font files (regular, bold) for a font name, or None."""
    try:
        import winreg
    except ImportError:          # browser / non-Windows
        return None

    found = {}
    windows_fonts = os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts")
    for hive in (winreg.HKEY_LOCAL_MACHINE, winreg.HKEY_CURRENT_USER):
        try:
            key = winreg.OpenKey(hive, r"SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts")
        except OSError:
            continue
        i = 0
        while True:
            try:
                label, path, _ = winreg.EnumValue(key, i)
            except OSError:
                break
            i += 1
            label = re.sub(r"\s*\((TrueType|OpenType)\)$", "", label).strip().lower()
            if not path.lower().endswith(".ttf"):
                continue
            if label == name:
                found.setdefault("", os.path.join(windows_fonts, path))
            elif label == name + " bold":
                found.setdefault("B", os.path.join(windows_fonts, path))

    if "" not in found:
        return None
    return found[""], found.get("B")


# =========================================================
# HELPERS
# =========================================================
def _colour(colour, default=None):
    if colour is None:
        return default
    if colour.type == "rgb" and isinstance(colour.rgb, str) and len(colour.rgb) >= 6:
        h = colour.rgb[-6:]
        return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))
    if colour.type == "theme":
        return THEME_COLOURS.get(colour.theme, default)
    return default


def _col_points(width):
    # Excel character width -> pixels -> points
    return (width * 7 + 5) * 0.75


def _group(digits, indian):
    if not indian or len(digits) <= 3:
        return f"{int(digits):,}"
    head, tail = digits[:-3], digits[-3:]
    parts = []
    while len(head) > 2:
        parts.insert(0, head[-2:])
        head = head[:-2]
    if head:
        parts.insert(0, head)
    return ",".join(parts + [tail])


def format_value(value, number_format):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime.datetime):
        return value.strftime("%d-%m-%Y")
    if isinstance(value, datetime.date):
        return value.strftime("%d-%m-%Y")
    if not isinstance(value, (int, float)):
        return str(value)

    nf = number_format or "General"

    if nf == "General":
        if isinstance(value, float) and not value.is_integer():
            return f"{value:.10g}"
        return str(int(value))

    if is_date_format(nf):
        try:
            return from_excel(value).strftime("%d-%m-%Y")
        except Exception:
            return str(value)

    section = nf.split(";")[0]
    section = re.sub(r'"[^"]*"|\[[^\]]*\]|\\.|_.|\*.', "", section)

    percent = "%" in section
    m = re.search(r"\.([0#]+)", section)
    decimals = len(m.group(1)) if m else 0
    indian = "##,##," in section
    grouped = "," in section

    v = value * 100 if percent else value
    text = f"{abs(v):.{decimals}f}"

    if grouped:
        whole, _, frac = text.partition(".")
        text = _group(whole, indian) + ("." + frac if frac else "")

    if v < 0 and float(text.replace(",", "")) != 0:
        text = "-" + text
    if percent:
        text += "%"
    return text


def _clean(text):
    """Normalise line breaks/tabs and drop other control characters."""
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\t", " ")
    return "".join(
        ch for ch in text
        if ch == "\n" or unicodedata.category(ch) != "Cc"
    )


class _Pdf(FPDF):
    """FPDF that picks a font per character, so mixed scripts never fail."""

    def setup_fonts(self, font_dir):
        self.font_dir = font_dir
        self.coverage = {}
        self.metrics = {}
        self.fake_bold = set()
        self.run_scale = 1.0
        self.family_of = {}
        self.run_family = DEFAULT_FAMILY
        for family in FALLBACKS + [DEFAULT_FAMILY]:
            self._register(family, *[os.path.join(font_dir, f) for f in BUNDLED[family]])

        # Gujarati/Hindi need "shaping" (vowel signs, conjuncts):
        # - desktop: uharfbuzz via fpdf2 (works through cell(), not text())
        # - browser: no uharfbuzz in Pyodide, so the page's canvas draws
        #   those runs (window.renderIndic in app.js) as crisp images
        self.shaping, self.canvas, self.canvas_cache = False, None, {}
        try:
            self.set_text_shaping(True)
            self.shaping = True
        except Exception:
            try:
                import js
                self.canvas = getattr(js, "renderIndic", None)
            except ImportError:
                pass

    def footer(self):
        """Small 'Page n of N' at the bottom of every page."""
        text = f"Page {self.page} of {getattr(self, 'total_pages', self.page)}"
        with self.local_context():
            self.set_font(DEFAULT_FAMILY, "", 8)
            self.set_text_color(120, 120, 120)
            self.text((A4_W - self.get_string_width(text)) / 2, A4_H - 14, text)

    def _register(self, family, regular, bold):
        if family not in self.coverage:
            if bold is None:
                # no bold file installed: bold is drawn with an outline instead
                self.fake_bold.add(family)
                bold = regular
            self.add_font(family, "", regular)
            self.add_font(family, "B", bold)
            self.coverage[family] = set(self.fonts[family.lower()].cmap)
            for style, path in (("", regular), ("B", bold)):
                tt = TTFont(path, lazy=True)
                self.metrics[(family, style)] = (
                    tt.getBestCmap(), tt["hmtx"].metrics, tt["head"].unitsPerEm
                )

    def _excel_width(self, family, part):
        """
        Width Excel gives this text: it lays out at 96 dpi with a whole-pixel
        font size and whole-pixel character widths, which is a few percent
        narrower than the font's exact width.
        """
        cmap, hmtx, upm = self.metrics[(family, self.run_style)]
        base_size = self.run_size / self.run_scale
        ppem = max(1, round(base_size * 96 / 72))
        px = 0
        for ch in part:
            glyph = cmap.get(ord(ch))
            if glyph is not None:
                px += round(hmtx[glyph][0] * ppem / upm)
        return px * 0.75 * self.run_scale

    def family_for(self, font_name):
        """Same font as Excel on Windows, else a font with identical widths."""
        name = (font_name or "").strip().lower()
        if name not in self.family_of:
            files = _system_font_files(name) if name else None
            if files:
                family = "sys-" + re.sub(r"\W", "", name)
                self._register(family, *files)
            else:
                family = METRIC_FONTS.get(name, DEFAULT_FAMILY)
                self._register(family, *[os.path.join(self.font_dir, f) for f in BUNDLED[family]])
            self.family_of[name] = family
        return self.family_of[name]

    def use(self, style, size, family=None, scale=None):
        self.run_style, self.run_size = style, size
        if family:
            self.run_family = family
        if scale:
            self.run_scale = scale

    def _render_canvas(self, family, part):
        key = (family, self.run_style, round(self.run_size, 2), part)
        if key not in self.canvas_cache:
            r, g, b = self.text_rgb
            res = self.canvas(part, family, self.run_style == "B", self.run_size, r, g, b)
            self.canvas_cache[key] = (
                res.png.to_bytes(), res.width, res.imageWidth, res.height, res.ascent
            )
        return self.canvas_cache[key]

    def _complex(self, family):
        return family in COMPLEX

    def runs(self, text):
        """Split text into [family, text] runs; unknown characters become '?'."""
        order = [self.run_family] + FALLBACKS
        result, family = [], None
        for ch in text:
            cp = ord(ch)
            if family is None or cp not in self.coverage[family]:
                family = next((f for f in order if cp in self.coverage[f]), None)
                if family is None:
                    family, ch = self.run_family, "?"
            if result and result[-1][0] == family:
                result[-1][1] += ch
            else:
                result.append([family, ch])
        return result

    def _run_width(self, family, part):
        if self.canvas is not None and self._complex(family):
            return self._render_canvas(family, part)[1]
        if not self._complex(family):
            return self._excel_width(family, part)
        self.set_font(family, self.run_style, self.run_size)
        return self.get_string_width(part)

    def measure(self, text):
        return sum(self._run_width(f, part) for f, part in self.runs(text))

    def draw(self, x, y, text):
        """Draw text with its baseline at y."""
        for family, part in self.runs(text):
            width = self._run_width(family, part)

            if self.canvas is not None and self._complex(family):
                png, _, image_w, h, ascent = self._render_canvas(family, part)
                self.image(BytesIO(png), x, y - ascent, image_w, h)

            elif self.shaping and self._complex(family):
                # fpdf2 puts the baseline at y + 0.5h + 0.3 * font size
                size = self.run_size
                self.set_xy(x, y - 0.8 * size)
                self.c_margin = 0
                self.cell(w=width, h=size, text=part)

            else:
                self.set_font(family, self.run_style, self.run_size)
                true_width = self.get_string_width(part)
                if not self._complex(family) and true_width > 0:
                    # squeeze to Excel's width so text breaks/cuts where Excel's does
                    self.set_stretching(100 * width / true_width)
                fake_bold = self.run_style == "B" and family in self.fake_bold
                if fake_bold:
                    self.set_draw_color(*self.text_rgb)
                    self.set_line_width(self.run_size * 0.08)
                    self.text_mode = "FILL_STROKE"
                self.text(x, y, part)
                if fake_bold:
                    self.text_mode = "FILL"
                self.set_stretching(100)

            x += width


def _wrap(pdf, text, width):
    lines = []
    for para in text.split("\n"):
        line = ""
        for word in para.split(" "):
            trial = f"{line} {word}" if line else word
            if line and pdf.measure(trial) > width:
                lines.append(line)
                line = word
            else:
                line = trial
        lines.append(line)
    return lines


# =========================================================
# SHEET RENDERER
# =========================================================
class _Sheet:
    def __init__(self, ws, pdf):
        self.ws = ws
        self.max_row = ws.max_row
        self.max_col = ws.max_column

        self.col_w = {}
        for c in range(1, self.max_col + 1):
            dim = ws.column_dimensions.get(get_column_letter(c))
            if dim is not None and dim.hidden:
                self.col_w[c] = 0
            else:
                width = dim.width if dim is not None and dim.width else DEFAULT_COL_WIDTH
                self.col_w[c] = _col_points(width)

        self.merge_of = {}       # (r, c) -> (min_row, min_col, max_row, max_col)
        for rng in ws.merged_cells.ranges:
            box = (rng.min_row, rng.min_col, rng.max_row, rng.max_col)
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    self.merge_of[(r, c)] = box

        default_h = ws.sheet_format.defaultRowHeight or DEFAULT_ROW_HEIGHT
        self.row_h = {}
        for r in range(1, self.max_row + 1):
            dim = ws.row_dimensions.get(r)
            if dim is not None and dim.hidden:
                self.row_h[r] = 0
            elif dim is not None and dim.height:
                self.row_h[r] = dim.height
            else:
                self.row_h[r] = self._autofit(pdf, r, default_h)

    def _autofit(self, pdf, r, default_h):
        """Row height Excel would give an automatic-height row."""
        height = default_h
        for c in range(1, self.max_col + 1):
            if (r, c) in self.merge_of:
                continue
            cell = self.ws.cell(r, c)
            text = _clean(format_value(cell.value, cell.number_format))
            if not text:
                continue
            size = cell.font.sz if cell.font is not None and cell.font.sz else 11
            lines = 1
            if cell.alignment is not None and cell.alignment.wrap_text:
                _set_font(pdf, cell.font, 1)
                lines = len(_wrap(pdf, text, self.col_w[c] - 2 * PAD))
            height = max(height, lines * size * LINE_RATIO)
        return height

    @property
    def width(self):
        return sum(self.col_w.values())

    def rows_height(self, rows):
        return sum(self.row_h[r] for r in rows)


def _set_font(pdf, font, scale):
    style = "B" if font is not None and font.b else ""
    size = max((font.sz if font is not None and font.sz else 11) * scale, 1)
    pdf.use(style, size, pdf.family_for(font.name if font is not None else None), scale)
    u = font is not None and font.u not in (None, "none")
    rgb = _colour(font.color if font is not None else None, (0, 0, 0))
    pdf.set_text_color(*rgb)
    pdf.text_rgb = rgb
    return size, u, rgb


def _is_empty(sheet, r, c):
    return (r, c) not in sheet.merge_of and sheet.ws.cell(r, c).value in (None, "")


def _draw_rows(pdf, sheet, rows, x0, y0, scale, links):
    ws = sheet.ws

    xs, x = {}, x0
    for c in range(1, sheet.max_col + 1):
        xs[c] = x
        x += sheet.col_w[c] * scale

    ys, y = {}, y0
    for r in rows:
        ys[r] = y
        y += sheet.row_h[r] * scale

    row_set = set(rows)

    def box_of(r, c):
        r1, c1, r2, c2 = sheet.merge_of.get((r, c), (r, c, r, c))
        w = sum(sheet.col_w[k] for k in range(c1, c2 + 1)) * scale
        h = sum(sheet.row_h[k] for k in range(r1, r2 + 1) if k in row_set) * scale
        return r1, c1, w, h

    # ---------------- fills ----------------
    for r in rows:
        if not sheet.row_h[r]:
            continue
        for c in range(1, sheet.max_col + 1):
            cell = ws.cell(r, c)
            if not cell.has_style or cell.fill is None or cell.fill.fill_type != "solid":
                continue
            rgb = _colour(cell.fill.fgColor)
            if rgb is None:
                continue
            pdf.set_fill_color(*rgb)
            pdf.rect(xs[c], ys[r], sheet.col_w[c] * scale, sheet.row_h[r] * scale, style="F")

    # ---------------- text ----------------
    for r in rows:
        for c in range(1, sheet.max_col + 1):
            r1, c1, w, h = box_of(r, c)
            if (r1, c1) != (r, c) or w <= 0 or h <= 0:
                continue

            cell = ws.cell(r, c)
            text = _clean(format_value(cell.value, cell.number_format))
            if not text:
                continue

            size, underline, rgb = _set_font(pdf, cell.font, scale)
            pad = PAD * scale
            al = cell.alignment
            horizontal = al.horizontal if al is not None else None
            vertical = (al.vertical if al is not None else None) or "bottom"

            is_num = isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool)
            if horizontal in (None, "general"):
                horizontal = "right" if is_num else "left"

            x, y = xs[c], ys[r]
            wrap = al is not None and al.wrap_text
            lines = _wrap(pdf, text, w - 2 * pad) if wrap else text.split("\n")[:1]

            # numbers never get cut: shrink them to fit the column
            if is_num and not wrap and w > 2 * pad:
                tw = pdf.measure(lines[0])
                if tw > w - 2 * pad:
                    size *= (w - 2 * pad) / tw
                    pdf.use(pdf.run_style, size)

            line_h = size * LINE_RATIO
            block_h = line_h * len(lines)
            if vertical == "top":
                ty = y
            elif vertical in ("center", "justify", "distributed"):
                ty = y + (h - block_h) / 2
            else:
                ty = y + h - block_h

            # Like Excel: unwrapped text may run into EMPTY cells beside it,
            # but is always cut at the first filled cell and at the row edges.
            clip_x1, clip_x2 = x, x + w
            if not wrap and not is_num:
                c_last = sheet.merge_of.get((r, c), (r, c, r, c))[3]
                if horizontal in ("left", "center", "centerContinuous", "fill"):
                    k = c_last + 1
                    while k <= sheet.max_col and _is_empty(sheet, r, k):
                        clip_x2 += sheet.col_w[k] * scale
                        k += 1
                if horizontal in ("right", "center", "centerContinuous"):
                    k = c - 1
                    while k >= 1 and _is_empty(sheet, r, k):
                        clip_x1 -= sheet.col_w[k] * scale
                        k -= 1

            def draw():
                for i, line in enumerate(lines):
                    tw = pdf.measure(line)
                    if horizontal == "right":
                        tx = x + w - pad - tw
                    elif horizontal in ("center", "centerContinuous"):
                        tx = x + (w - tw) / 2
                    else:
                        tx = x + pad + (al.indent or 0) * 7 * 0.75 * scale
                    baseline = ty + line_h * i + size * 1.0
                    pdf.draw(tx, baseline, line)
                    if underline:
                        pdf.set_line_width(0.5 * scale)
                        pdf.set_draw_color(*rgb)
                        pdf.line(tx, baseline + size * 0.12, tx + tw, baseline + size * 0.12)

            # local_context keeps fpdf2's font state in step with the PDF's
            # once the clipping area is closed (otherwise fonts get mixed up)
            with pdf.local_context():
                with pdf.rect_clip(clip_x1, y, clip_x2 - clip_x1, h):
                    draw()

            target = cell.hyperlink.location if cell.hyperlink else None
            if not target and cell.hyperlink is not None and cell.hyperlink.target:
                target = cell.hyperlink.target
            if target:
                name = target.lstrip("#").split("!")[0].strip("'")
                if name in links:
                    pdf.link(x, y, w, h, links[name])

    # ---------------- borders ----------------
    def edge(style, colour, x1, y1, x2, y2):
        if not style:
            return
        pdf.set_line_width(BORDER_WIDTHS.get(style, 0.5) * max(scale, 0.5))
        pdf.set_draw_color(*(_colour(colour, (0, 0, 0))))
        dash = DASHES.get(style)
        if dash:
            pdf.set_dash_pattern(dash=dash[0], gap=dash[1])
        pdf.line(x1, y1, x2, y2)
        if dash:
            pdf.set_dash_pattern()

    for r in rows:
        if not sheet.row_h[r]:
            continue
        for c in range(1, sheet.max_col + 1):
            if not sheet.col_w[c]:
                continue
            cell = ws.cell(r, c)
            if not cell.has_style or cell.border is None:
                continue
            b = cell.border
            box = sheet.merge_of.get((r, c))
            x1, y1 = xs[c], ys[r]
            x2, y2 = x1 + sheet.col_w[c] * scale, y1 + sheet.row_h[r] * scale

            # skip edges that lie inside a merged block
            inside = lambda rr, cc: box is not None and sheet.merge_of.get((rr, cc)) == box
            if not inside(r - 1, c):
                edge(b.top.style, b.top.color, x1, y1, x2, y1)
            if not inside(r + 1, c):
                edge(b.bottom.style, b.bottom.color, x1, y2, x2, y2)
            if not inside(r, c - 1):
                edge(b.left.style, b.left.color, x1, y1, x1, y2)
            if not inside(r, c + 1):
                edge(b.right.style, b.right.color, x2, y1, x2, y2)


def _margins(ws):
    m = ws.page_margins
    return m.left * 72, m.right * 72, m.top * 72, m.bottom * 72


def _render_one_page(pdf, ws, links):
    sheet = _Sheet(ws, pdf)
    rows = list(range(1, sheet.max_row + 1))

    left, right, top, bottom = _margins(ws)
    avail_w = A4_W - left - right
    avail_h = A4_H - top - bottom

    grid_w, grid_h = sheet.width, sheet.rows_height(rows)
    scale = min(1.0, avail_w / grid_w if grid_w else 1, avail_h / grid_h if grid_h else 1)

    x0 = left + (avail_w - grid_w * scale) / 2
    pdf.add_page()
    _draw_rows(pdf, sheet, rows, x0, top, scale, links)


# =========================================================
# CONTENTS PAGE(S)
# =========================================================
# (title, width in points, alignment) - widths add up to the A4 text width
CONTENTS_COLUMNS = [
    ("Sr. No.", 55, "C"),
    ("City Name", 255, "L"),
    ("Page No.", 70, "C"),
    ("Remarks", A4_W - 2 * contents.PAGE_MARGIN_X - 380, "L"),
]
HEADER_RGB = (31, 78, 121)
GRID_RGB = (191, 199, 209)
STRIPE_RGB = (242, 245, 249)
ZERO_RGB = (255, 199, 206)
ZERO_TEXT_RGB = (156, 0, 6)
LINK_RGB = (5, 99, 193)


def _contents_cell(pdf, text, x, y, w, h, align, style, size, rgb):
    pdf.use(style, size, DEFAULT_FAMILY, 1)
    pdf.set_text_color(*rgb)
    pdf.text_rgb = rgb

    pad = 6
    tw = pdf.measure(text)
    if align == "C":
        tx = x + (w - tw) / 2
    else:
        tx = x + pad

    with pdf.local_context():
        with pdf.rect_clip(x, y, w, h):
            pdf.draw(tx, y + h / 2 + size * 0.36, text)


def _draw_contents(pdf, entries, links):
    """entries: list of (city name, page number, remark)."""
    x0 = contents.PAGE_MARGIN_X
    per_page = contents.ROWS_PER_PAGE
    chunks = [entries[i:i + per_page] for i in range(0, len(entries), per_page)] or [[]]

    for n, chunk in enumerate(chunks):
        pdf.add_page()
        y = contents.PAGE_TOP

        # title
        title = "Contents" if n == 0 else "Contents (continued)"
        _contents_cell(pdf, title, x0, y, 400, 26, "L", "B", 18, HEADER_RGB)
        pdf.set_draw_color(*HEADER_RGB)
        pdf.set_line_width(1.2)
        pdf.line(x0, y + 32, A4_W - x0, y + 32)
        y += contents.TITLE_H

        # header row
        x = x0
        for name, width, _ in CONTENTS_COLUMNS:
            pdf.set_fill_color(*HEADER_RGB)
            pdf.rect(x, y, width, contents.HEADER_H, style="F")
            _contents_cell(pdf, name, x, y, width, contents.HEADER_H,
                           "C", "B", 10, (255, 255, 255))
            x += width
        y += contents.HEADER_H

        # city rows
        for i, (city, page, remark) in enumerate(chunk):
            sr = n * per_page + i + 1
            zero = bool(remark)
            fill = ZERO_RGB if zero else (STRIPE_RGB if i % 2 else None)

            x = x0
            values = [str(sr), city, str(page), remark]
            for col, ((_, width, align), value) in enumerate(zip(CONTENTS_COLUMNS, values)):
                if fill:
                    pdf.set_fill_color(*fill)
                    pdf.rect(x, y, width, contents.ROW_H, style="F")
                pdf.set_draw_color(*GRID_RGB)
                pdf.set_line_width(0.5)
                pdf.rect(x, y, width, contents.ROW_H)

                if col == 1:
                    style, rgb = "", LINK_RGB
                elif col == 3:
                    style, rgb = "B", ZERO_TEXT_RGB
                else:
                    style, rgb = "", (0, 0, 0)
                _contents_cell(pdf, value, x, y, width, contents.ROW_H, align, style, 10, rgb)
                x += width

            pdf.link(x0, y, x - x0, contents.ROW_H, links[city])
            y += contents.ROW_H


# =========================================================
# MAIN ENTRY
# =========================================================
def build_pdf(wb, font_dir, title="SLB Gazette"):
    pdf = _Pdf(orientation="P", unit="pt", format="A4")
    pdf.set_auto_page_break(False)
    pdf.set_title(title)
    pdf.setup_fonts(font_dir)

    cities = [ws for ws in wb.worksheets if ws.title.strip().lower() != "index"]
    remarks = {}
    for ws in wb.worksheets:
        if ws.title.strip().lower() == "index":
            remarks = dict(contents.read_index(ws))

    # Contents first, then one page per city (same numbering as the Excel Index)
    pages = contents.city_page_numbers(len(cities))
    pdf.total_pages = contents.contents_page_count(len(cities)) + len(cities)
    links = {ws.title: pdf.add_link(page=p) for ws, p in zip(cities, pages)}

    _draw_contents(
        pdf,
        [(ws.title, p, remarks.get(ws.title, "")) for ws, p in zip(cities, pages)],
        links,
    )

    for ws in cities:
        _render_one_page(pdf, ws, links)

    return bytes(pdf.output())

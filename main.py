import os
import sys
from copy import copy
from tkinter import Tk, filedialog

import win32com.client as win32
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

# PDF renderer + fonts are shared with the web version
WEB_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "web")
sys.path.insert(0, WEB_DIR)

from contents import write_index
from pdf_report import build_pdf

# =========================================================
# CONFIG
# =========================================================
DELETE_START_ROW = 40
DELETE_ROW_COUNT = 5
DEFAULT_EXCEL_WIDTH = 17.6

OUTPUT_NAME = "SLB Gazette file.xlsx"


def is_all_zero(ws):
    """True when the sheet has numbers and every one of them is 0."""
    found_number = False

    for row in ws.iter_rows(values_only=True):
        for value in row:
            if isinstance(value, bool) or not isinstance(value, (int, float)):
                continue
            if value != 0:
                return False
            found_number = True

    return found_number


# =========================================================
# SELECT INPUT FOLDER
# =========================================================
Tk().withdraw()

input_folder = filedialog.askdirectory(
    title="Select Folder Containing XLS Files"
)

if not input_folder:
    print("No folder selected.")
    exit()

# =========================================================
# FIND XLS FILES
# =========================================================
excel_files = sorted([
    f for f in os.listdir(input_folder)
    if f.endswith(".xls") and not f.startswith("~$")
])

if not excel_files:
    print("No XLS files found.")
    exit()

# =========================================================
# SHOW FILES
# =========================================================
print("\nAvailable Files:")
for i, f in enumerate(excel_files, start=1):
    print(f"{i}. {f}")

print("\nEnter file numbers in required merge order")
print("Example: 3,1,2")

order_input = input("Order: ")

try:
    order_indexes = [int(x.strip()) - 1 for x in order_input.split(",")]
    sorted_files = [excel_files[i] for i in order_indexes]
except:
    print("Invalid order.")
    exit()

# =========================================================
# CONVERT XLS TO XLSX
# =========================================================
temp_folder = os.path.join(input_folder, "converted_xlsx")
os.makedirs(temp_folder, exist_ok=True)

converted_files = []

print("\nConverting XLS to XLSX...")

excel = win32.Dispatch("Excel.Application")
excel.Visible = False
excel.DisplayAlerts = False

for file_name in sorted_files:

    try:
        xls_path = os.path.abspath(
            os.path.join(input_folder, file_name)
        )

        xlsx_name = os.path.splitext(file_name)[0] + ".xlsx"

        xlsx_path = os.path.abspath(
            os.path.join(temp_folder, xlsx_name)
        )

        wb = excel.Workbooks.Open(xls_path)
        wb.SaveAs(xlsx_path, FileFormat=51)
        wb.Close(False)

        converted_files.append(xlsx_path)

        print(f"Converted: {xlsx_name}")

    except Exception as e:
        print(f"\nERROR converting: {file_name}")
        print(e)

excel.Quit()

# =========================================================
# CREATE OUTPUT WORKBOOK
# =========================================================
merged_wb = Workbook()
merged_wb.remove(merged_wb.active)

# =========================================================
# CREATE INDEX SHEET
# =========================================================
# Filled in after all sheets are added (see write_index below)
index_ws = merged_wb.create_sheet(title="Index", index=0)

index_entries = []

# =========================================================
# PROCESS EACH FILE
# =========================================================
for file_path in converted_files:

    print(f"\nProcessing: {os.path.basename(file_path)}")

    source_wb = load_workbook(file_path)

    for sheet_name in source_wb.sheetnames:

        source_ws = source_wb[sheet_name]

        # -------------------------------------------------
        # UNIQUE SHEET NAME
        # -------------------------------------------------
        new_sheet_name = sheet_name
        counter = 1

        while new_sheet_name in merged_wb.sheetnames:
            new_sheet_name = f"{sheet_name}_{counter}"
            counter += 1

        target_ws = merged_wb.create_sheet(title=new_sheet_name)

        # -------------------------------------------------
        # COPY CELL VALUES + STYLES
        # -------------------------------------------------
        for row in source_ws.iter_rows():

            for cell in row:

                new_cell = target_ws[cell.coordinate]

                new_cell.value = cell.value

                if cell.has_style:
                    new_cell.font = copy(cell.font)
                    new_cell.border = copy(cell.border)
                    new_cell.fill = copy(cell.fill)
                    new_cell.number_format = copy(cell.number_format)
                    new_cell.protection = copy(cell.protection)
                    new_cell.alignment = copy(cell.alignment)

        # -------------------------------------------------
        # COPY COLUMN WIDTHS
        # -------------------------------------------------
        for col_idx in range(1, source_ws.max_column + 1):

            col_letter = get_column_letter(col_idx)

            src_dim = source_ws.column_dimensions.get(col_letter)

            if src_dim and src_dim.width:
                target_ws.column_dimensions[col_letter].width = src_dim.width
            else:
                target_ws.column_dimensions[col_letter].width = DEFAULT_EXCEL_WIDTH

        # -------------------------------------------------
        # FORCE D & E SAME AS C
        # -------------------------------------------------
        c_width = target_ws.column_dimensions["C"].width

        target_ws.column_dimensions["D"].width = c_width
        target_ws.column_dimensions["E"].width = c_width

        # -------------------------------------------------
        # COPY ROW HEIGHTS
        # -------------------------------------------------
        for r, dim in source_ws.row_dimensions.items():
            target_ws.row_dimensions[r].height = dim.height

        # -------------------------------------------------
        # COPY MERGED CELLS
        # -------------------------------------------------
        for merged_cell in source_ws.merged_cells.ranges:
            target_ws.merge_cells(str(merged_cell))

        # -------------------------------------------------
        # DELETE ROWS 40-44
        # -------------------------------------------------
        delete_end_row = (
            DELETE_START_ROW + DELETE_ROW_COUNT - 1
        )

        # Unmerge overlapping merged cells
        for merged in list(target_ws.merged_cells.ranges):

            if (
                merged.min_row <= delete_end_row
                and merged.max_row >= DELETE_START_ROW
            ):
                target_ws.unmerge_cells(str(merged))

        # Delete rows
        target_ws.delete_rows(
            DELETE_START_ROW,
            DELETE_ROW_COUNT
        )

        # Remove leftover row formatting
        for r in range(
            DELETE_START_ROW,
            DELETE_START_ROW + DELETE_ROW_COUNT
        ):
            if r in target_ws.row_dimensions:
                del target_ws.row_dimensions[r]

        # -------------------------------------------------
        # INDEX ENTRY (highlighted when every number is 0)
        # -------------------------------------------------
        all_zero = is_all_zero(target_ws)

        index_entries.append((new_sheet_name, all_zero))

        print(f"  Added Sheet: {new_sheet_name}"
              + ("  [ALL ZERO]" if all_zero else ""))

# =========================================================
# WRITE INDEX: Sr. No. | City Name | Page No. | Remarks
# =========================================================
write_index(index_ws, index_entries)

# =========================================================
# SAVE INITIAL FILE
# =========================================================
output_file = os.path.join(input_folder, OUTPUT_NAME)

merged_wb.save(output_file)

# =========================================================
# REOPEN FOR FINAL PROCESSING
# =========================================================
wb = load_workbook(output_file)

# =========================================================
# MERGE LAST ROW A:E
# =========================================================
for ws in wb.worksheets:

    if ws.title.strip().lower() == "index":
        print(f"Skipped sheet: {ws.title}")
        continue

    # -----------------------------------------------------
    # FIND LAST NON-EMPTY ROW
    # -----------------------------------------------------
    last_row = ws.max_row

    while last_row > 1:

        row_has_data = any(
            ws.cell(last_row, col).value is not None
            for col in range(1, ws.max_column + 1)
        )

        if row_has_data:
            break

        last_row -= 1

    print(f"{ws.title} -> Last Row: {last_row}")

    # -----------------------------------------------------
    # MERGE A:E IN LAST ROW
    # -----------------------------------------------------
    merge_range = f"A{last_row}:E{last_row}"

    ws.merge_cells(merge_range)

    ws[f"A{last_row}"].alignment = Alignment(
        horizontal="left",
        vertical="center"
    )
    
    # -----------------------------------------------------
    # DELETE EXTRA ROWS AFTER LAST ROW
    # -----------------------------------------------------
    rows_after = ws.max_row - last_row

    if rows_after > 0:
        ws.delete_rows(last_row + 1, rows_after)

    # -----------------------------------------------------
    # SET PRINT AREA ONLY TILL LAST USED ROW
    # -----------------------------------------------------
    last_col_letter = get_column_letter(ws.max_column)

    ws.print_area = f"A1:{last_col_letter}{last_row}"

    # -----------------------------------------------------
    # PAGE SETUP
    # -----------------------------------------------------
    ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
    ws.page_setup.paperSize = ws.PAPERSIZE_A4

    # Fit content properly
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1

    # Enable fit-to-page
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # Proper margins for A4 PDF
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.4
    ws.page_margins.bottom = 0.4
    ws.page_margins.header = 0.2
    ws.page_margins.footer = 0.2

    # Center horizontally
    ws.print_options.horizontalCentered = True
    ws.print_options.verticalCentered = False

    # Repeat first row while printing (optional)
    ws.print_title_rows = "1:1"

    # Set print area
    last_col_letter = get_column_letter(ws.max_column)
    ws.print_area = f"A1:{last_col_letter}{last_row}"


# =========================================================
# SAVE FINAL FILE
# =========================================================
base, ext = os.path.splitext(output_file)

final_output = f"{base}_merged{ext}"

wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True

wb.save(final_output)

# =========================================================
# EXPORT PDF REPORT (A4, ONE CITY PER PAGE)
# =========================================================
# Drawn directly (not via Excel) so the page is always A4,
# whatever printer is installed on this PC.
pdf_output = f"{base}_merged.pdf"

print("\nCreating PDF...")

try:
    with open(pdf_output, "wb") as f:
        f.write(build_pdf(wb, os.path.join(WEB_DIR, "fonts")))

    print(f"PDF saved: {pdf_output}")

except Exception as e:
    pdf_output = None
    print("\nERROR creating PDF")
    print(e)

# =========================================================
# COMPLETE
# =========================================================
print("\n===================================")
print("Merge Completed Successfully")
print("Saved File:")
print(final_output)
if pdf_output:
    print(pdf_output)
print("===================================")